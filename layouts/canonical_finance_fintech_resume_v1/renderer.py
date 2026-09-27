"""canonical_finance_fintech_resume_v1 renderer.

FROZEN layout (see LAYOUT_SPEC.yaml). This module renders ANY resume content
with the exact approved Alpaca V4.1 visual layout. Agents must NOT change
layout constants here; only the `content` dict varies per role profile.

Content schema:
{
  "name": str,
  "contact_line": str,
  "portfolio_links": [{"text": str, "url": str}, ...],  # rendered as "Portfolio: <l1> | GitHub: <l2>"
  "sections": [
    {"kind": "bullets", "heading": str, "bullets": [str, ...]},
    {"kind": "entries", "heading": str, "entries": [
        {"title": str, "location": str, "date": str, "role": str | None, "bullets": [str, ...]}
    ]},
    {"kind": "lines", "heading": str, "lines": [str, ...]},
    {"kind": "single_bold_line", "heading": str, "text": str},
    {"kind": "labeled_line", "label": str, "text": str},  # e.g. LANGUAGES footer
  ]
}
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# --- Frozen layout constants (mirror LAYOUT_SPEC.yaml) ---
NAVY = RGBColor(0x1F, 0x38, 0x64)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x40, 0x40, 0x40)
FONT = "Calibri"
BODY_PT = 10


def _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=2,
          left_indent=None, hanging=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if left_indent is not None:
        pf.left_indent = Inches(left_indent)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    return p


def _run(p, text, size=BODY_PT, bold=False, italic=False, color=INK,
         underline=False):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    if underline:
        r.font.underline = True
    return r


def _hyperlink(p, text, url, size=9.5):
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:sz", str(int(size * 2))),
                     ("w:color", "0563C1"), ("w:u", "single")):
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), FONT)
            el.set(qn("w:hAnsi"), FONT)
        else:
            el.set(qn("w:val"), val)
        rPr.append(el)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hl.append(new_run)
    p._p.append(hl)


def _rule_bottom(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single",
        qn("w:sz"): "4",
        qn("w:color"): "1F3864",
        qn("w:space"): "1",
    })
    pBdr.append(bottom)
    pPr.append(pBdr)


def _heading(doc, text):
    p = _para(doc, before=6, after=2)
    _run(p, text.upper(), size=10.5, bold=True, color=NAVY)
    _rule_bottom(p)


def _bullet(doc, text):
    p = _para(doc, before=0, after=2, left_indent=0.3, hanging=0.16)
    _run(p, "- ", size=BODY_PT, color=INK)
    _run(p, text, size=BODY_PT, color=INK)


def _entry(doc, title, location, date):
    p = _para(doc, before=4, after=0)
    _run(p, title, size=BODY_PT, bold=True, color=INK)
    _run(p, f" | {location} | ", size=BODY_PT, color=INK)
    _run(p, date, size=BODY_PT, bold=True, color=INK)


def _role_line(doc, title):
    p = _para(doc, after=1)
    _run(p, title, size=BODY_PT, bold=True, color=NAVY)


def render_canonical_resume(content: dict[str, Any], out_docx: str | Path) -> Path:
    """Render `content` with the frozen canonical layout. Returns DOCX path."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Header
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    _run(p, content["name"], size=17, bold=True, color=NAVY)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    _run(p, content["contact_line"], size=9.5, color=MUTED)
    links = content.get("portfolio_links") or []
    if links:
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
        _run(p, "Portfolio: ", size=9.5, color=INK)
        _hyperlink(p, links[0]["text"], links[0]["url"])
        if len(links) > 1:
            _run(p, " | GitHub: ", size=9.5, color=INK)
            _hyperlink(p, links[1]["text"], links[1]["url"])

    for sec in content["sections"]:
        kind = sec["kind"]
        if kind == "bullets":
            _heading(doc, sec["heading"])
            for b in sec["bullets"]:
                _bullet(doc, b)
        elif kind == "entries":
            _heading(doc, sec["heading"])
            for e in sec["entries"]:
                _entry(doc, e["title"], e["location"], e["date"])
                if e.get("role"):
                    _role_line(doc, e["role"])
                for b in e["bullets"]:
                    _bullet(doc, b)
        elif kind == "lines":
            _heading(doc, sec["heading"])
            for i, line in enumerate(sec["lines"]):
                p = _para(doc, after=1 if i < len(sec["lines"]) - 1 else 2)
                _run(p, line, size=BODY_PT, color=INK)
        elif kind == "single_bold_line":
            _heading(doc, sec["heading"])
            p = _para(doc, after=2)
            _run(p, sec["text"], size=BODY_PT, bold=True, color=INK)
        elif kind == "labeled_line":
            p = _para(doc, before=2, after=2)
            _run(p, sec["label"] + "  ", size=BODY_PT, bold=True, color=NAVY)
            _run(p, sec["text"], size=BODY_PT, color=INK)
        else:
            raise ValueError(f"unknown section kind: {kind}")

    out = Path(out_docx)
    doc.save(str(out))
    return out
